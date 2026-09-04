# -*- coding: utf-8 -*-
"""
GitHub Daily Radar - Excel Exporter Module
Xuất và lưu trữ lũy tiến dữ liệu dự án vào file Excel (.xlsx).
"""

import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def get_default_excel_path():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, "reports", "github_radar_archive.xlsx")

def append_or_update_excel(categories_data, excel_path=None):
    """
    Lưu hoặc cập nhật danh sách repository vào file Excel lũy tiến.
    Nếu repo đã tồn tại, cập nhật số sao và ngày quét gần nhất.
    Nếu là repo mới, thêm dòng mới vào bảng.
    """
    if excel_path is None:
        excel_path = get_default_excel_path()
        
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    today_str = datetime.now().strftime("%Y-%m-%d")

    # Mở file hiện có hoặc tạo file mới
    if os.path.exists(excel_path):
        try:
            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active
        except Exception:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "GitHub Radar Archive"
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "GitHub Radar Archive"

    headers = [
        "Ngày quét", 
        "Danh mục", 
        "Tên Repository", 
        "Stars", 
        "Forks", 
        "Ngôn ngữ", 
        "Mô tả", 
        "Tags / Topics", 
        "Link GitHub"
    ]

    # Thiết lập Header nếu sheet trống
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        ws.append(headers)
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

    # Lập chỉ mục các repo đã có sẵn (để tránh trùng lặp)
    # key: repo_full_name -> row_index
    existing_repos = {}
    for row in range(2, ws.max_row + 1):
        repo_name = ws.cell(row=row, column=3).value
        if repo_name:
            existing_repos[repo_name.strip()] = row

    new_count = 0
    updated_count = 0

    body_font = Font(name="Segoe UI", size=10)
    link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    for cat in categories_data:
        cat_name = cat["name"]
        for r in cat["repos"]:
            full_name = r["full_name"].strip()
            tags_str = ", ".join(r["topics"])
            
            if full_name in existing_repos:
                # Cập nhật số liệu mới nhất
                row_idx = existing_repos[full_name]
                ws.cell(row=row_idx, column=1, value=today_str)
                ws.cell(row=row_idx, column=4, value=r["stars"])
                ws.cell(row=row_idx, column=5, value=r["forks"])
                ws.cell(row=row_idx, column=7, value=r["description"])
                updated_count += 1
            else:
                # Thêm dòng mới
                new_row = [
                    today_str,
                    cat_name,
                    full_name,
                    r["stars"],
                    r["forks"],
                    r["language"],
                    r["description"],
                    tags_str,
                    r["url"]
                ]
                ws.append(new_row)
                current_row = ws.max_row
                existing_repos[full_name] = current_row
                
                # Format dòng mới
                for c in range(1, len(new_row) + 1):
                    cell = ws.cell(row=current_row, column=c)
                    cell.font = body_font
                    if c in [1, 4, 5, 6]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(vertical="center")
                        
                    if current_row % 2 == 0:
                        cell.fill = alt_fill
                        
                # Gán hyperlink cho cột Link GitHub
                link_cell = ws.cell(row=current_row, column=9)
                link_cell.hyperlink = r["url"]
                link_cell.font = link_font
                
                new_count += 1

    # Tự động căn chỉnh độ rộng cột
    col_widths = {1: 14, 2: 30, 3: 28, 4: 12, 5: 12, 6: 15, 7: 50, 8: 30, 9: 40}
    for col_num, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width

    wb.save(excel_path)
    print(f"[+] Đã lưu vào Excel: {excel_path} (+{new_count} mới, {updated_count} cập nhật, Tổng {len(existing_repos)} dự án trong kho)")
    return excel_path
